"""Auditoría de listas externas: ¿el email es de la empresa donde la persona trabaja HOY?

Pensado para bases antiguas validadas solo por entregabilidad (ZeroBounce, NeverBounce...):
un buzón 'valid' de la empresa anterior no es un contacto válido.
"""
import csv
from collections import Counter
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .almacen import Dominios, norm_empresa
from .emails import _dominio, dominio_coincide, motivo_rechazo

NO_VISIBLE = {"", "(no visible)", "no visible", "-", "?"}
EXCLUIR = ("jubilad", "retired", "in remembrance", "en memoria", "fallecid")
COLORES = {"APTO": "C6EFCE", "REVISAR": "FFEB9C", "PENDIENTE": "EDEDED", "NO APTO": "FFC7CE"}


def _col(cabeceras: list[str], *candidatas: str) -> str | None:
    bajas = {c.lower().strip(): c for c in cabeceras}
    for cand in candidatas:
        if cand.lower() in bajas:
            return bajas[cand.lower()]
    return None


def leer_tabla(ruta: Path) -> list[dict]:
    if ruta.suffix.lower() in (".xlsx", ".xlsm"):
        ws = load_workbook(ruta, read_only=True, data_only=True).active
        filas = list(ws.iter_rows(values_only=True))
        cab = [str(c or "").strip() for c in filas[0]]
        return [{cab[i]: ("" if v is None else str(v)) for i, v in enumerate(f) if i < len(cab)} for f in filas[1:]]
    texto = ruta.read_text(encoding="utf-8-sig")
    sep = ";" if texto.split("\n", 1)[0].count(";") >= texto.split("\n", 1)[0].count(",") else ","
    return list(csv.DictReader(texto.splitlines(), delimiter=sep))


def misma_empresa(a: str, b: str) -> bool:
    na, nb = norm_empresa(a), norm_empresa(b)
    if not na or not nb:
        return False
    return na == nb or na in nb or nb in na or na.split()[0] == nb.split()[0]


def veredicto(fila: dict, c_email: str, c_actual: str | None, c_origen: str | None,
              c_titular: str | None, dominios: Dominios | None = None) -> tuple[str, str]:
    """Misma regla que el sistema: solo es APTO un email del dominio de la empresa actual de LinkedIn."""
    email = (fila.get(c_email) or "").strip()
    dom = _dominio(email)
    actual = (fila.get(c_actual) or "").strip() if c_actual else ""
    origen = (fila.get(c_origen) or "").strip() if c_origen else ""
    titular = (fila.get(c_titular) or "").lower() if c_titular else ""
    if not dom:
        return "NO APTO", "Sin email"
    if any(p in titular for p in EXCLUIR):
        return "NO APTO", "Persona jubilada o fallecida según LinkedIn"
    if actual.lower() in NO_VISIBLE:
        return "PENDIENTE", "No se ha verificado la empresa actual en LinkedIn: no usar hasta verificarla"
    verificados = dominios.verificados(actual) if dominios else []
    motivo = motivo_rechazo(email, actual, verificados)
    if not motivo:
        return "APTO", f"El dominio {dom} corresponde a la empresa actual ({actual})"
    if origen and not misma_empresa(origen, actual) and dominio_coincide(email, origen):
        motivo += f"; es de la empresa anterior ({origen})"
    return "NO APTO", motivo[0].upper() + motivo[1:]


def auditar(ruta: Path, salida: Path) -> tuple[Path, Counter]:
    filas = leer_tabla(ruta)
    if not filas:
        raise SystemExit(f"{ruta} está vacío")
    cab = list(filas[0].keys())
    c_email = _col(cab, "Email", "E-mail", "Correo", "email")
    if not c_email:
        raise SystemExit(f"No encuentro la columna de email en {ruta.name}")
    c_actual = _col(cab, "Empresa actual", "Empresa LinkedIn", "Current company") or _col(cab, "Empresa", "Company")
    c_origen = _col(cab, "Empresa", "Company") if c_actual != _col(cab, "Empresa", "Company") else None
    c_titular = _col(cab, "Titular actual", "Titular / cargo actual (LinkedIn)", "Cargo actual")

    wb = Workbook()
    ws = wb.active
    ws.title = "Auditoría"
    ws.append(["Veredicto", "Motivo"] + cab)
    stats: Counter = Counter()
    dominios = Dominios()
    for f in filas:
        v, motivo = veredicto(f, c_email, c_actual, c_origen, c_titular, dominios)
        stats[v] += 1
        ws.append([v, motivo] + [f.get(c, "") for c in cab])
        ws.cell(ws.max_row, 1).fill = PatternFill("solid", fgColor=COLORES[v])
    for c in ws[1]:
        c.font = Font(bold=True)
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cab) + 2)}{ws.max_row}"
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 60

    rs = wb.create_sheet("Resumen")
    rs.append(["Archivo", ruta.name])
    rs.append(["Columna email", c_email])
    rs.append(["Columna empresa actual", c_actual or "-"])
    rs.append(["Columna empresa de origen", c_origen or "-"])
    rs.append([])
    for v in ("APTO", "PENDIENTE", "NO APTO"):
        rs.append([v, stats[v]])
    rs.append([])
    rs.append(["Regla", "Solo es APTO un email cuyo dominio es el de la empresa actual en LinkedIn. "
                        "Que el buzón exista (ZeroBounce 'valid') no basta."])
    rs.column_dimensions["A"].width = 28

    salida.parent.mkdir(parents=True, exist_ok=True)
    wb.save(salida)
    return salida, stats
