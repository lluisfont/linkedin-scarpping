"""Exportación de los leads a una hoja Excel con formato."""
from collections import Counter
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

COLUMNAS = [  # (cabecera, campo, ancho)
    ("Nombre", "nombre", 26), ("Segmento", "segmentos", 12), ("Cargo", "cargo", 38), ("Empresa", "empresa", 28),
    ("Ubicación", "ubicacion", 24), ("Email", "email", 34), ("Estado email", "email_estado", 13),
    ("Score", "email_score", 8), ("Fuente email", "email_fuente", 11), ("Dominio", "dominio", 20),
    ("LinkedIn", "url_linkedin", 42), ("Sales Navigator", "url_salesnav", 20),
    ("Tiempo en puesto", "tiempo_puesto", 16), ("Búsqueda", "busqueda", 22),
    ("Primera vez visto", "fecha_alta", 16), ("Última vez visto", "ultima_vez", 16),
    ("Resultado email", "resultado_email", 16),
    ("Validación", "validacion", 12), ("Nota", "nota", 40),
    ("Email descartado (no es de la empresa actual)", "email_descartado", 34),
]
COLORES_ESTADO = {
    "valid": "C6EFCE", "accept_all": "FFEB9C", "webmail": "FFEB9C",
    "unknown": "EDEDED", "invalid": "FFC7CE", "disposable": "FFC7CE",
}
CABECERA = PatternFill("solid", fgColor="1F4E78")


def _hoja_tabla(ws, cabeceras: list[str], filas: list[list], nombre_tabla: str, anchos=None):
    ws.append(cabeceras)
    for fila in filas:
        ws.append(fila)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = CABECERA
        c.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"
    for i, ancho in enumerate(anchos or [18] * len(cabeceras), 1):
        ws.column_dimensions[get_column_letter(i)].width = ancho
    if filas:
        ref = f"A1:{get_column_letter(len(cabeceras))}{len(filas) + 1}"
        tabla = Table(displayName=nombre_tabla, ref=ref)
        tabla.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
        ws.add_table(tabla)


def exportar(leads: list[dict], busquedas: list[dict], ruta: Path) -> Path:
    wb = Workbook()

    # --- Prospectos ---
    ws = wb.active
    ws.title = "Prospectos"
    filas = []
    for l in leads:
        fila = [l.get(campo, "") for _, campo, _ in COLUMNAS]
        i_score = [c for _, c, _ in COLUMNAS].index("email_score")
        fila[i_score] = int(fila[i_score]) if str(fila[i_score]).isdigit() else None
        filas.append(fila)
    _hoja_tabla(ws, [c for c, _, _ in COLUMNAS], filas, "Prospectos", [a for _, _, a in COLUMNAS])
    campos = [c for _, c, _ in COLUMNAS]
    col_estado = campos.index("email_estado") + 1
    for fila in ws.iter_rows(min_row=2):
        color = COLORES_ESTADO.get(str(fila[col_estado - 1].value or ""))
        if color:
            fila[col_estado - 1].fill = PatternFill("solid", fgColor=color)
        for campo in ("url_linkedin", "url_salesnav"):
            celda = fila[campos.index(campo)]
            if celda.value:
                celda.hyperlink = str(celda.value)
                celda.font = Font(color="0563C1", underline="single")

    # --- Resumen ---
    rs = wb.create_sheet("Resumen")
    total = len(leads)
    con_email = sum(1 for l in leads if l.get("email"))
    rs.append(["Generado", datetime.now().strftime("%Y-%m-%d %H:%M")])
    rs.append(["Contactos (personas únicas)", total])
    rs.append(["Con email", con_email])
    rs.append(["% con email", f"{(100 * con_email / total):.0f}%" if total else "-"])
    rs.append(["Emails descartados (dominio de otra empresa)", sum(1 for l in leads if l.get("email_descartado"))])
    rs.append([])
    rs.append(["Estado del email", "Contactos"])
    for estado, n in Counter(l.get("email_estado") or "(sin email)" for l in leads).most_common():
        rs.append([estado, n])
    for titulo, campo in (("Segmento", "segmentos"), ("Búsqueda", "busqueda")):
        rs.append([])
        rs.append([titulo, "Contactos", "Con email"])
        total_c: Counter = Counter()
        con_e: Counter = Counter()
        for l in leads:
            for b in (l.get(campo) or "(sin etiqueta)").split(" | "):
                total_c[b] += 1
                con_e[b] += bool(l.get("email"))
        for b, n in total_c.most_common():
            rs.append([b, n, con_e[b]])
    for fila in rs.iter_rows():
        if fila[0].value in ("Estado del email", "Búsqueda", "Segmento"):
            for c in fila:
                c.font = Font(bold=True)
    rs.column_dimensions["A"].width = 34
    rs.column_dimensions["B"].width = 14
    rs.column_dimensions["C"].width = 12

    # --- Búsquedas ---
    bs = wb.create_sheet("Búsquedas")
    cab = ["Fecha", "Segmento", "Búsqueda", "Origen", "Página", "Total resultados", "Contactos en lote",
           "Nuevos en la base", "URL"]
    filas_b = [[b.get("fecha"), b.get("segmento"), b.get("busqueda"), b.get("origen"), b.get("pagina"),
                b.get("total_resultados"), b.get("leads_en_lote"), b.get("nuevos"), b.get("url")]
               for b in busquedas]
    _hoja_tabla(bs, cab, filas_b, "Busquedas", [16, 12, 24, 16, 8, 16, 16, 16, 60])

    ruta.parent.mkdir(parents=True, exist_ok=True)
    try:
        wb.save(ruta)
    except PermissionError:  # el archivo está abierto en Excel
        ruta = ruta.with_name(f"{ruta.stem}_{datetime.now():%Y%m%d_%H%M%S}{ruta.suffix}")
        wb.save(ruta)
    return ruta
