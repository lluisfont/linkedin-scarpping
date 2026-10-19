import json
import tempfile
import unittest
from pathlib import Path
from unittest import mock

from openpyxl import load_workbook

from prospector import almacen
from prospector.emails import buscar_email_lead, enriquecer
from prospector.excel import exportar, tamano_empresa
from prospector.hunter import Hunter, SinCreditos
from prospector.nombres import candidatos, es_incompleto, limpiar, partir
from prospector.prospeo import Prospeo


class Resp:
    def __init__(self, status, cuerpo):
        self.status_code, self._cuerpo = status, cuerpo
        self.content = b"x"
        self.text = json.dumps(cuerpo)

    def json(self):
        return self._cuerpo


class FakeSession:
    """Devuelve respuestas según una función (url, params|json) -> Resp y registra las llamadas."""

    def __init__(self, fn):
        self.fn, self.llamadas = fn, []

    def get(self, url, params=None, timeout=None):
        self.llamadas.append((url, dict(params or {})))
        return self.fn(url, params)

    def post(self, url, json=None, headers=None, timeout=None):
        self.llamadas.append((url, json))
        return self.fn(url, json)


class TestNombres(unittest.TestCase):
    def test_limpiar(self):
        self.assertEqual(limpiar("Ana Pérez García, MBA 🚀"), "Ana Pérez García")
        self.assertEqual(limpiar("Dr. Juan López (He/Him)"), "Juan López")
        self.assertEqual(limpiar("Marta Ruiz | HR Director"), "Marta Ruiz")

    def test_partir(self):
        self.assertEqual(partir("Ana García López"), ("Ana", ["García", "López"]))
        self.assertEqual(partir("María José de la Fuente Ruiz"), ("María José", ["de la Fuente", "Ruiz"]))
        self.assertEqual(partir("Juan Carlos"), ("Juan", ["Carlos"]))

    def test_candidatos(self):
        self.assertEqual(candidatos("Juan Carlos Pérez Gómez"), [
            ("Juan Carlos", "Pérez"), ("Juan", "Pérez"), ("Juan Carlos", "Pérez Gómez")])
        self.assertEqual(candidatos("Laura Martí"), [("Laura", "Martí")])
        self.assertEqual(candidatos("Ana M."), [])

    def test_incompleto(self):
        self.assertTrue(es_incompleto("Ana M."))
        self.assertFalse(es_incompleto("Ana Martí"))


class TestAlmacen(unittest.TestCase):
    def test_fusionar_y_dedupe(self):
        leads = []
        a = almacen.normalizar_lead({"id": "ACwAA1", "nombre": "Ana Martí", "cargo": "HR Director",
                                     "empresa": "Seat S.A."}, "rrhh")
        n, _, _ = almacen.fusionar(leads, [a])
        self.assertEqual(n, 1)
        self.assertEqual(leads[0]["url_linkedin"], "https://www.linkedin.com/in/ACwAA1")
        # mismo lead desde un perfil público: se fusiona por nombre+empresa y gana la URL personalizada
        b = almacen.normalizar_lead({"url": "https://www.linkedin.com/in/ana-marti/", "nombre": "Ana Martí",
                                     "empresa": "SEAT SA", "cargo": "Chief People Officer"}, "perfiles")
        n, act, _ = almacen.fusionar(leads, [b])
        self.assertEqual((n, act, len(leads)), (0, 1, 1))
        self.assertEqual(leads[0]["url_linkedin"], "https://www.linkedin.com/in/ana-marti")
        self.assertEqual(leads[0]["cargo"], "Chief People Officer")
        self.assertEqual(leads[0]["busqueda"], "rrhh | perfiles")

    def test_fusionar_otra_empresa(self):
        def lead(empresa, eid, cargo="HR Director"):
            return almacen.normalizar_lead({"id": "ACwAA9", "nombre": "Pau Roca", "cargo": cargo,
                                            "empresa": empresa, "empresa_id": eid}, "rrhh-bcn-D")
        # con email: se queda el puesto del email y se marca para revisar
        leads = [lead("Affle", "66741")]
        leads[0].update(email="pau.roca@affle.com", resultado_email="encontrado")
        almacen.fusionar(leads, [lead("Jampp", "2483441", "Head of People")])
        self.assertEqual((leads[0]["empresa"], leads[0]["cargo"]), ("Affle", "HR Director"))
        self.assertEqual(leads[0]["validacion"], "Revisar")
        self.assertIn("también figura en Jampp", leads[0]["nota"])
        # sin email: pasa a la empresa nueva y se vuelve a buscar
        leads = [lead("Affle", "66741")]
        leads[0].update(resultado_email="no_encontrado", dominio="affle.com")
        almacen.fusionar(leads, [lead("Jampp", "2483441")])
        self.assertEqual((leads[0]["empresa"], leads[0]["resultado_email"], leads[0]["dominio"]), ("Jampp", "", ""))
        # mismo ID de empresa con otro nombre visible: no es un cambio
        leads = [lead("INDCRESA", "5367409")]
        leads[0].update(email="a@indcresa.com")
        almacen.fusionar(leads, [lead("INDCRESA / Cacao en polvo desde 1915", "5367409")])
        self.assertEqual(leads[0]["validacion"], "")

    def test_descarta_miembro_anonimo(self):
        self.assertIsNone(almacen.normalizar_lead({"id": "x", "nombre": "LinkedIn Member"}, "b"))

    def test_base_sqlite_ida_y_vuelta(self):
        with tempfile.TemporaryDirectory() as d:
            db = Path(d) / "contactos.db"
            lead = almacen.normalizar_lead({"id": "1", "nombre": "Pep Solé", "cargo": "Dir; RRHH"}, "rrhh-bcn")
            almacen.guardar_leads([lead], db)
            fila = almacen.cargar_leads(db)[0]
            self.assertEqual((fila["cargo"], fila["segmentos"]), ("Dir; RRHH", "RRHH"))

    def test_misma_persona_en_rrhh_y_ceo_no_se_duplica(self):
        """Hoy RRHH, mañana CEOs: una sola ficha por persona con los dos segmentos y su email conservado."""
        with tempfile.TemporaryDirectory() as d:
            db = Path(d) / "contactos.db"
            leads = []
            a = almacen.normalizar_lead({"id": "ACwAA7", "nombre": "Marta Ruiz", "cargo": "HR Director",
                                         "empresa": "Acme"}, "rrhh-bcn-1000")
            _, _, ids = almacen.fusionar(leads, [a])
            leads[0]["email"] = "marta.ruiz@acme.com"
            almacen.guardar_leads(leads, db)
            almacen.registrar_apariciones(ids, "rrhh-bcn-1000", "RRHH", db)

            leads = almacen.cargar_leads(db)
            b = almacen.normalizar_lead({"id": "ACwAA7", "nombre": "Marta Ruiz", "cargo": "CEO",
                                         "empresa": "Acme"}, "ceo-bcn")
            nuevos, existian, ids = almacen.fusionar(leads, [b])
            almacen.guardar_leads(leads, db)
            almacen.registrar_apariciones(ids, "ceo-bcn", "CEO", db)

            base = almacen.cargar_leads(db)
            self.assertEqual((nuevos, existian, len(base)), (0, 1, 1))
            self.assertEqual(base[0]["segmentos"], "RRHH | CEO")
            self.assertEqual(base[0]["cargo"], "CEO")                      # dato actualizado
            self.assertEqual(base[0]["email"], "marta.ruiz@acme.com")      # no se vuelve a pagar
            from prospector.emails import pendientes
            self.assertEqual(pendientes(base), [])
            with almacen.closing(almacen.conectar(db)) as con:
                self.assertEqual(con.execute("SELECT COUNT(*) FROM apariciones").fetchone()[0], 2)

    def test_segmento_de_etiqueta(self):
        self.assertEqual(almacen.segmento_de("rrhh-bcn-1000"), "RRHH")
        self.assertEqual(almacen.segmento_de("ceo-madrid"), "CEO")

    def test_lote_txt_de_trozos(self):
        with tempfile.TemporaryDirectory() as d:
            ruta = Path(d) / "p1.txt"
            ruta.write_text("# pagina: 1\n# total: 247\nACwAAA1|Ana Martí|HR Director|Seat|248021|Barcelona|9 months\nFIN\n",
                            encoding="utf-8")
            meta, leads = almacen.leer_lote(ruta)
            self.assertEqual((meta["pagina"], meta["total"]), ("1", "247"))
            self.assertEqual(leads[0]["empresa_id"], "248021")
            self.assertTrue(leads[0]["url_salesnav"].startswith("https://www.linkedin.com/sales/lead/ACwAAA1"))

    def test_dominio(self):
        self.assertEqual(almacen.limpiar_dominio("https://www.Seat.es/es/"), "seat.es")
        self.assertEqual(almacen.norm_empresa("Banco Sabadell, S.A."), "banco sabadell")


def _hunter(fn):
    s = FakeSession(fn)
    return Hunter("k", session=s, pausa=0), s


class TestEmails(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dominios = almacen.Dominios(Path(self.tmp.name) / "dominios.json")

    def tearDown(self):
        self.tmp.cleanup()

    def _lead(self, nombre="Juan Carlos Pérez Gómez", empresa="Acme S.L."):
        return almacen.normalizar_lead({"id": "ACwAA9", "nombre": nombre, "empresa": empresa}, "b")

    def test_prueba_combinaciones_y_cachea_dominio(self):
        def fn(url, p):
            if p.get("first_name") == "Juan":
                return Resp(200, {"data": {"email": "juan.perez@acme.com", "score": 91, "domain": "acme.com",
                                           "verification": {"status": "valid"}}})
            return Resp(200, {"data": {"email": None, "domain": "acme.com"}})
        hunter, s = _hunter(fn)
        lead = self._lead()
        self.assertEqual(buscar_email_lead(lead, hunter, None, self.dominios), "encontrado")
        self.assertEqual((lead["email"], lead["email_score"], lead["email_estado"]),
                         ("juan.perez@acme.com", "91", "valid"))
        # 1.ª llamada por empresa; la 2.ª ya usa el dominio aprendido
        self.assertEqual(s.llamadas[0][1].get("company"), "Acme S.L.")
        self.assertEqual(s.llamadas[1][1].get("domain"), "acme.com")
        self.assertEqual(self.dominios.get("ACME SL"), "acme.com")

    def test_regla_empresa_actual(self):
        from prospector.emails import motivo_rechazo
        # caso Villalonga: buzón 'valid' de la empresa ANTERIOR -> rechazado siempre
        self.assertIn("empresa actual", motivo_rechazo("jv@lacaixa.es", "Longtown HR"))
        self.assertIn("empresa actual", motivo_rechazo("s@meetproducerx.com", "Mandarin Oriental Hotel Group"))
        self.assertIn("otro país", motivo_rechazo("a@culligan.cz", "Culligan International"))
        self.assertIn("personal", motivo_rechazo("a@icab.cat", "Sanofi"))
        self.assertIn("sin empresa", motivo_rechazo("a@acme.com", ""))
        self.assertEqual(motivo_rechazo("t@thegbfoods.com", "GBfoods"), "")
        # siglas
        for email, empresa in (("a@pg.com", "Procter & Gamble"), ("b@db.com", "Deutsche Bank SAE"),
                               ("c@fbg.ub.es", "Fundació Bosch i Gimpera"), ("e@df-sa.es", "DF Electric, S.A."),
                               ("d@barcelona.msf.org", "Médicos Sin Fronteras España")):
            self.assertEqual(motivo_rechazo(email, empresa), "", email)
        # un dominio distinto al nombre solo vale si está verificado (página de la empresa en LinkedIn)
        empresa = "Freudenberg Home & Cleaning Solucions"
        self.assertNotEqual(motivo_rechazo("b@fhp-ww.com", empresa), "")
        self.assertEqual(motivo_rechazo("b@fhp-ww.com", empresa, ["fhp-ww.com"]), "")

    def test_email_de_otra_empresa_se_descarta(self):
        hunter, s = _hunter(lambda u, p: Resp(200, {"data": {
            "email": "silvia@meetproducerx.com", "score": 83, "domain": "meetproducerx.com",
            "verification": {"status": "valid"}}}))
        lead = self._lead("Sílvia Forés Sabater", "Mandarin Oriental Hotel Group")
        self.assertEqual(buscar_email_lead(lead, hunter, None, self.dominios), "descartado")
        self.assertEqual(lead["email"], "")
        self.assertEqual(lead["email_descartado"], "silvia@meetproducerx.com")
        self.assertIn("Descartado:", lead["nota"])
        # el dominio equivocado no se aprende para otros contactos de la empresa
        self.assertEqual(self.dominios.get("Mandarin Oriental Hotel Group"), "")
        self.assertEqual(len(s.llamadas), 1)  # no insiste con otras variantes del nombre

    def test_dominio_verificado_tiene_prioridad(self):
        self.dominios.set("Acme S.L.", "acme-group.com", verificado=True)
        self.dominios.set("Acme S.L.", "acme.cz")  # una propuesta de Hunter no pisa el verificado
        hunter, s = _hunter(lambda u, p: Resp(200, {"data": {"email": "laura.marti@acme-group.com",
                                                             "verification": {"status": "valid"}}}))
        lead = self._lead("Laura Martí")
        self.assertEqual(buscar_email_lead(lead, hunter, None, self.dominios), "encontrado")
        self.assertEqual(s.llamadas[0][1].get("domain"), "acme-group.com")

    def test_451_excluye(self):
        hunter, _ = _hunter(lambda u, p: Resp(451, {"errors": [{"id": "claimed_email"}]}))
        lead = self._lead("Laura Martí")
        self.assertEqual(buscar_email_lead(lead, hunter, None, self.dominios), "excluido_451")

    def test_prospeo_respaldo(self):
        hunter, _ = _hunter(lambda u, p: Resp(200, {"data": {"email": None}}))
        sp = FakeSession(lambda u, j: Resp(200, {"error": False, "person": {"email": {
            "email": "laura@acme.com", "status": "VERIFIED", "revealed": True}}, "company": {"domain": "acme.com"}}))
        prospeo = Prospeo("k", session=sp, pausa=0)
        lead = self._lead("Laura Martí")
        self.assertEqual(buscar_email_lead(lead, hunter, prospeo, self.dominios), "encontrado")
        self.assertEqual((lead["email"], lead["email_fuente"], lead["email_estado"]),
                         ("laura@acme.com", "prospeo", "valid"))
        self.assertEqual(sp.llamadas[0][1]["data"]["company_name"], "Acme S.L.")

    def test_sin_creditos_para_y_guarda(self):
        hunter, _ = _hunter(lambda u, p: Resp(429, {"errors": [{"id": "too_many_requests"}]}))
        leads = [self._lead("Laura Martí")]
        guardado = mock.Mock()
        with self.assertRaises(SinCreditos):
            enriquecer(leads, hunter=hunter, prospeo=None, dominios=self.dominios,
                       guardar=guardado, log=lambda *_: None)
        guardado.assert_called()

    def test_no_reintenta_no_encontrados(self):
        hunter, s = _hunter(lambda u, p: Resp(200, {"data": {"email": None}}))
        leads = [self._lead("Laura Martí")]
        enriquecer(leads, hunter=hunter, prospeo=None, dominios=self.dominios, log=lambda *_: None)
        self.assertEqual(leads[0]["resultado_email"], "no_encontrado")
        n = len(s.llamadas)
        enriquecer(leads, hunter=hunter, prospeo=None, dominios=self.dominios, log=lambda *_: None)
        self.assertEqual(len(s.llamadas), n)


class TestExcel(unittest.TestCase):
    def test_exportar(self):
        lead = almacen.normalizar_lead({"id": "ACwAA1", "nombre": "Ana Martí", "empresa": "Seat"}, "rrhh")
        lead.update(email="ana@seat.es", email_estado="valid", email_score="95")
        with tempfile.TemporaryDirectory() as d:
            ruta = exportar([lead], [{"busqueda": "rrhh", "url": "u"}], Path(d) / "p.xlsx")
            wb = load_workbook(ruta)
            self.assertEqual(wb.sheetnames, ["Prospectos", "Resumen", "Búsquedas"])
            ws = wb["Prospectos"]
            col = {c.value: c.column_letter for c in ws[1]}
            self.assertEqual(ws[f"{col['Email']}2"].value, "ana@seat.es")
            self.assertEqual(ws[f"{col['Score']}2"].value, 95)
            self.assertEqual(ws[f"{col['Segmento']}2"].value, "RRHH")
            self.assertEqual(ws[f"{col['LinkedIn']}2"].hyperlink.target, "https://www.linkedin.com/in/ACwAA1")

    def test_tamano_empresa(self):
        self.assertEqual(tamano_empresa("rrhh-bcn-D"), "51-200 (puede ser <100)")
        self.assertEqual(tamano_empresa("rrhh-bcn-I"), "10000+")
        self.assertEqual(tamano_empresa("rrhh"), "")
        self.assertEqual(tamano_empresa("rrhh-bcn-201 | rrhh-bcn-E | rrhh-bcn-D"), "201-500")
        self.assertEqual(tamano_empresa("rrhh-bcn-201 | rrhh-bcn-D"), "201+")

    def test_limpiar_abreviatura_maria(self):
        self.assertEqual(limpiar("Mª Luisa Gambús Sáiz"), "María Luisa Gambús Sáiz")


if __name__ == "__main__":
    unittest.main()
